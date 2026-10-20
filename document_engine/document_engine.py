
import argparse
import json
import re
import sys
from pathlib import Path

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
    ListFlowable,
    ListItem,
)

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor


def safe_name(value):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(value)).strip("_") or "vva"


def as_list(value):
    if isinstance(value, list):
        return [str(item) for item in value]
    if value:
        return [str(value)]
    return []


def load_json(path):
    with Path(path).open("r", encoding="utf-8") as file:
        data = json.load(file)
    if not isinstance(data, dict):
        raise ValueError(f"Expected a JSON object in {path}")
    return data


def make_sections(ai, project):
    return [
        ("Executive Summary", [ai.get("summary", "Not provided.")]),
        ("Purpose and Rationale", [ai.get("purposeAndRationale", "Not provided.")]),
        ("Changed Components", as_list(ai.get("changedComponents"))),
        ("Technical Changes", as_list(ai.get("technicalChanges"))),
        ("Technologies", as_list(ai.get("technologies"))),
        ("Architectural Impact", [ai.get("architecturalImpact", "Not provided.")]),
        ("Observations and Limitations", as_list(ai.get("observationsAndLimitations"))),
        ("Commit Information", [
            f"Repository: {project.get('repository', 'Unknown')}",
            f"Branch: {project.get('branch', 'Unknown')}",
            f"Commit: {project.get('analyzedCommit', 'Unknown')}",
            f"Commit message: {project.get('commitMessage', 'Not provided')}",
        ]),
    ]


def create_pdf(ai, project, path):
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="VVASection",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#176B68"),
        spaceBefore=12,
        spaceAfter=6,
    ))
    styles.add(ParagraphStyle(
        name="VVABody",
        parent=styles["BodyText"],
        alignment=TA_LEFT,
        leading=15,
        spaceAfter=6,
    ))

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(20 * mm, 12 * mm, "VVA — Version & Workflow Automator")
        canvas.drawRightString(190 * mm, 12 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=20 * mm,
        leftMargin=20 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=str(ai.get("title", "VVA Technical Report")),
    )

    story = [
        Paragraph("VVA Technical Report", styles["Title"]),
        Spacer(1, 6 * mm),
        Paragraph(str(ai.get("title", "Untitled Analysis")), styles["Heading1"]),
        Spacer(1, 4 * mm),
    ]

    for heading, items in make_sections(ai, project):
        story.append(Paragraph(heading, styles["VVASection"]))
        if not items:
            items = ["No details provided."]
        if len(items) == 1:
            story.append(Paragraph(items[0], styles["VVABody"]))
        else:
            bullets = [
                ListItem(Paragraph(item, styles["VVABody"]))
                for item in items
            ]
            story.append(ListFlowable(bullets, bulletType="bullet"))
        story.append(Spacer(1, 2 * mm))

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def create_docx(ai, project, path):
    doc = Document()
    doc.add_heading("VVA Technical Synopsis", level=0)
    doc.add_heading(str(ai.get("title", "Untitled Analysis")), level=1)

    for heading, items in make_sections(ai, project):
        doc.add_heading(heading, level=2)
        if not items:
            doc.add_paragraph("No details provided.")
        elif len(items) == 1:
            doc.add_paragraph(items[0])
        else:
            for item in items:
                doc.add_paragraph(item, style="List Bullet")

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    doc.save(str(path))


def create_pptx(ai, project, path):
    prs = Presentation()
    prs.core_properties.title = str(ai.get("title", "VVA Presentation"))

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = str(ai.get("title", "VVA Project Report"))
    slide.placeholders[1].text = (
        f"Repository: {project.get('repository', 'Unknown')}\n"
        f"Commit: {project.get('analyzedCommit', 'Unknown')}"
    )

    sections = make_sections(ai, project)

    for heading, items in sections:
        if not items:
            items = ["No details provided."]

        # Split longer lists across multiple slides.
        chunks = [items[i:i + 5] for i in range(0, len(items), 5)]

        for chunk in chunks:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = heading
            frame = slide.placeholders[1].text_frame
            frame.clear()

            for index, item in enumerate(chunk):
                paragraph = (
                    frame.paragraphs[0]
                    if index == 0
                    else frame.add_paragraph()
                )
                paragraph.text = item
                paragraph.font.size = PptPt(20)
                paragraph.font.color.rgb = PptRGBColor(45, 55, 72)

    prs.save(str(path))


def main():
    parser = argparse.ArgumentParser(description="VVA document generation engine")
    parser.add_argument("--ai-input", required=True)
    parser.add_argument("--project-input", required=True)
    parser.add_argument("--output-dir", required=True)
    parser.add_argument("--format", choices=["pdf", "docx", "pptx", "all"], default="all")
    args = parser.parse_args()

    ai = load_json(args.ai_input)
    project = load_json(args.project_input)

    output_dir = Path(args.output_dir).resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    repository = safe_name(project.get("repository", "vva"))
    commit = safe_name(str(project.get("analyzedCommit", "unknown"))[:7])

    generators = {
        "pdf": (
            f"{repository}_{commit}_report.pdf",
            create_pdf,
        ),
        "docx": (
            f"{repository}_{commit}_synopsis.docx",
            create_docx,
        ),
        "pptx": (
            f"{repository}_{commit}_presentation.pptx",
            create_pptx,
        ),
    }

    formats = list(generators) if args.format == "all" else [args.format]
    results = []

    for fmt in formats:
        filename, generator = generators[fmt]
        output_path = output_dir / filename
        generator(ai, project, output_path)

        if not output_path.is_file() or output_path.stat().st_size == 0:
            raise RuntimeError(f"Document generation failed: {output_path}")

        results.append({
            "format": fmt,
            "filename": filename,
            "path": str(output_path),
            "sizeBytes": output_path.stat().st_size,
        })

    print(json.dumps(results))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"VVA Document Engine error: {exc}", file=sys.stderr)
        sys.exit(1)
